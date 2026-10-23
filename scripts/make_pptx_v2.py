#!/usr/bin/env python3
"""Build the Aegis evaluation presentation (professional dark deck)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG = RGBColor(0x0A, 0x0E, 0x17)
PANEL = RGBColor(0x14, 0x1C, 0x2E)
BORDER = RGBColor(0x2A, 0x3A, 0x5E)
INK = RGBColor(0xE7, 0xED, 0xF7)
MUTED = RGBColor(0x8F, 0xA0, 0xBD)
ACCENT = RGBColor(0x6E, 0xA0, 0xFF)
GREEN = RGBColor(0x2F, 0xD5, 0x7A)
RED = RGBColor(0xFF, 0x5D, 0x6C)
AMBER = RGBColor(0xFF, 0xB4, 0x54)
CYAN = RGBColor(0x38, 0xC8, 0xDD)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def rect(slide, x, y, w, h, color=PANEL, line=BORDER, radius=0.06):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = radius
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True, size=16, color=INK, bold=False, leading=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    first = True
    for content, style in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = style.get('align', align)
        p.space_after = Pt(style.get('after', 6))
        p.space_before = Pt(style.get('before', 0))
        p.line_spacing = style.get('leading', leading)
        r = p.add_run()
        r.text = content
        f = r.font
        f.size = Pt(style.get('size', size))
        f.bold = style.get('bold', bold)
        f.italic = style.get('italic', False)
        f.color.rgb = style.get('color', color)
        f.name = 'Segoe UI'
    return tb


def chip(slide, x, y, w, h, label, color):
    sp = rect(slide, x, y, w, h, color=color, radius=0.5)
    tf = sp.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x06, 0x12, 0x2E)
    r.font.name = 'Segoe UI'
    return sp


def header(slide, num, title):
    text(slide, Inches(0.7), Inches(0.55), Inches(9), Inches(0.9),
         title, size=30, bold=True)
    text(slide, Inches(11.9), Inches(0.6), Inches(0.9), Inches(0.6),
         [('0%d' % num, {'size': 12, 'color': MUTED, 'align': PP_ALIGN.RIGHT})])
    text(slide, Inches(0.7), Inches(1.25), Inches(12), Inches(0.3),
         [('Aegis  ·  Self-Healing Multi-Cloud Platform', {'size': 12, 'color': MUTED})])
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.05), Inches(1.6), Pt(2.5))
    ln.fill.solid()
    ln.fill.fore_color.rgb = ACCENT
    ln.line.fill.background()
    ln.shadow.inherit = False


def bullet(slide, x, y, w, h, items):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for content, style in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(style.get('after', 9))
        p.line_spacing = 1.08
        r1 = p.add_run()
        r1.text = '▸  '
        r1.font.size = Pt(style.get('size', 16))
        r1.font.color.rgb = style.get('color', ACCENT)
        r1.font.bold = True
        r1.font.name = 'Segoe UI'
        r2 = p.add_run()
        r2.text = content
        r2.font.size = Pt(style.get('size', 16))
        r2.font.color.rgb = style.get('color', INK)
        r2.font.bold = style.get('bold', False)
        r2.font.name = 'Segoe UI'
    return tb


def kpi(slide, x, y, w, h, label, value, sub, color=ACCENT):
    rect(slide, x, y, w, h)
    text(slide, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), Inches(0.3),
         [(label.upper(), {'size': 10.5, 'color': MUTED})])
    text(slide, x + Inches(0.25), y + Inches(0.5), w - Inches(0.5), Inches(0.6),
         [(value, {'size': 26, 'bold': True, 'color': color})])
    text(slide, x + Inches(0.25), y + Inches(1.15), w - Inches(0.5), Inches(0.3),
         [(sub, {'size': 11, 'color': MUTED})])


# ============ SLIDE 1 — TITLE ============
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, Inches(0), Inches(0), SW, SH, color=BG, line=None)
rect(s, Inches(0.9), Inches(2.15), Inches(11.5), Inches(3.3), color=PANEL, radius=0.03)
text(s, Inches(0.9), Inches(2.55), Inches(11.5), Inches(0.5),
     [('MAJOR PROJECT · FINAL EVALUATION', {'size': 13, 'color': ACCENT, 'align': PP_ALIGN.CENTER})])
text(s, Inches(0.9), Inches(3.05), Inches(11.5), Inches(1.1),
     [('Aegis', {'size': 60, 'bold': True, 'align': PP_ALIGN.CENTER})])
text(s, Inches(0.9), Inches(3.95), Inches(11.5), Inches(0.6),
     [('Self-Healing Multi-Cloud Platform', {'size': 24, 'color': INK, 'align': PP_ALIGN.CENTER})])
text(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(0.4),
     [('Run one app across Render · Vercel · Cloud Run  —  monitor, score, route, heal — at ₹0/month',
       {'size': 14, 'color': MUTED, 'align': PP_ALIGN.CENTER})])
text(s, Inches(0.9), Inches(6.25), Inches(11.5), Inches(0.4),
     [('Manav Sarvaiya  ·  B.E. Final Year  ·  2026-09', {'size': 13, 'color': MUTED, 'align': PP_ALIGN.CENTER})])

# ============ SLIDE 2 — PROBLEM & SOLUTION ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 2, 'The Problem → The Idea')
rect(s, Inches(0.7), Inches(2.0), Inches(5.8), Inches(4.7))
text(s, Inches(1.0), Inches(2.25), Inches(5.2), Inches(0.4),
     'The problem', size=17, bold=True, color=RED)
bullet(s, Inches(1.0), Inches(2.75), Inches(5.2), Inches(3.8), [
    ('One cloud = one point of failure.', {}),
    ('Free tiers are cheap but unreliable — Render sleeps after ~15 min idle.', {}),
    ('Commercial multi-cloud / DR tools are expensive and lock you in.', {}),
    ('No simple, ₹0 way to combine free tiers into something MORE reliable than one paid instance.', {}),
])
rect(s, Inches(6.7), Inches(2.0), Inches(5.9), Inches(4.7))
text(s, Inches(7.0), Inches(2.25), Inches(5.2), Inches(0.4),
     'The idea', size=17, bold=True, color=GREEN)
bullet(s, Inches(7.0), Inches(2.75), Inches(5.2), Inches(3.8), [
    ('Treat clouds as interchangeable, disposable resources.', {}),
    ('Probe every cloud every 5 min from OUTSIDE the fleet.', {}),
    ('Score each cloud 0–100 → route traffic to the healthiest.', {}),
    ('When one fails: log it, alert, and redeploy automatically — self-healing.', {}),
])

# ============ SLIDE 3 — ARCHITECTURE ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 3, 'Architecture (microservices)')
# control plane
rect(s, Inches(0.7), Inches(1.95), Inches(11.9), Inches(1.35), color=PANEL)
text(s, Inches(0.95), Inches(2.05), Inches(11.4), Inches(0.3),
     'CONTROL PLANE  (web · api · monitor · report — runs locally / Oracle VM)', size=11.5, bold=True, color=CYAN)
text(s, Inches(0.95), Inches(2.4), Inches(11.4), Inches(0.8),
     'web:3100  portal+realtime dashboard   ·   api:4000  keys+Swagger+rate-limit   ·   monitor  probe→score→heal→alert   ·   report  daily/weekly',
     size=12, color=MUTED)
# supabase
rect(s, Inches(0.7), Inches(3.6), Inches(11.9), Inches(0.85), color=PANEL)
text(s, Inches(0.95), Inches(3.72), Inches(11.4), Inches(0.3),
     'SUPABASE POSTGRES (free)', size=11.5, bold=True, color=ACCENT)
text(s, Inches(0.95), Inches(4.05), Inches(11.4), Inches(0.35),
     'health_checks · incidents · alerts · risk_flags · api_keys · audit_log · failover_events · health_reports',
     size=12, color=MUTED)
# vantage points
rect(s, Inches(0.7), Inches(4.75), Inches(3.8), Inches(0.85), color=PANEL)
text(s, Inches(0.95), Inches(4.9), Inches(3.3), Inches(0.55),
     [('Vantage 1\n', {'size': 11.5, 'bold': True, 'color': GREEN}), ('Cloudflare Worker · cron */5', {'size': 11, 'color': MUTED})])
rect(s, Inches(4.75), Inches(4.75), Inches(3.8), Inches(0.85), color=PANEL)
text(s, Inches(5.0), Inches(4.9), Inches(3.3), Inches(0.55),
     [('Vantage 2\n', {'size': 11.5, 'bold': True, 'color': GREEN}), ('GitHub Actions · cron */5', {'size': 11, 'color': MUTED})])
rect(s, Inches(8.8), Inches(4.75), Inches(3.8), Inches(0.85), color=PANEL)
text(s, Inches(9.05), Inches(4.9), Inches(3.3), Inches(0.55),
     [('Vantage 3\n', {'size': 11.5, 'bold': True, 'color': GREEN}), ('Oracle Always-Free VM · 24/7', {'size': 11, 'color': MUTED})])
# data plane
rect(s, Inches(0.7), Inches(5.9), Inches(11.9), Inches(0.85), color=PANEL, line=ACCENT)
text(s, Inches(0.95), Inches(6.02), Inches(11.4), Inches(0.3),
     'DATA PLANE (protected — same app, 3 clouds)', size=11.5, bold=True, color=AMBER)
text(s, Inches(0.95), Inches(6.35), Inches(11.4), Inches(0.35),
     'Render (free) ✅  ·  Vercel (Hobby) ✅  ·  Cloud Run (free tier) ➕   —   Cloudflare Worker routes to the healthiest',
     size=12, color=MUTED)

# ============ SLIDE 4 — HEALTH SCORING ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 4, 'How Aegis scores & routes')
rect(s, Inches(0.7), Inches(2.0), Inches(5.8), Inches(4.7))
text(s, Inches(1.0), Inches(2.25), Inches(5.2), Inches(0.4),
     'Health score (0–100)', size=17, bold=True)
bullet(s, Inches(1.0), Inches(2.75), Inches(5.2), Inches(1.6), [
    ('0.5 × uptime %', {'size': 18}),
    ('0.3 × latency score', {'size': 18}),
    ('0.2 × (1 − error rate)', {'size': 18}),
])
rect(s, Inches(1.0), Inches(4.65), Inches(5.2), Inches(1.8), color=BG, line=BORDER)
text(s, Inches(1.25), Inches(4.8), Inches(4.7), Inches(1.5),
     [('Grade bands\n', {'size': 12, 'color': MUTED}),
      ('≥90 Excellent   ·   ≥70 Good\n≥50 At risk   ·   <50 Critical', {'size': 14, 'bold': True})])
rect(s, Inches(6.7), Inches(2.0), Inches(5.9), Inches(4.7))
text(s, Inches(7.0), Inches(2.25), Inches(5.2), Inches(0.4),
     'Routing decision (deterministic)', size=17, bold=True)
bullet(s, Inches(7.0), Inches(2.75), Inches(5.2), Inches(2.2), [
    ('Rank clouds by health score (desc).', {}),
    ('Tie-break by lowest latency.', {}),
    ('Final tie-break by lowest carbon intensity (green routing).', {}),
])
rect(s, Inches(7.0), Inches(5.0), Inches(5.2), Inches(1.5), color=BG, line=BORDER)
text(s, Inches(7.25), Inches(5.15), Inches(4.7), Inches(1.2),
     [('Live right now (2026-09-17)\n', {'size': 11, 'color': MUTED}),
      ('vercel 93 → serve   ·   render 85 → standby', {'size': 14, 'bold': True, 'color': GREEN})])

# ============ SLIDE 5 — SELF-HEALING LOOP ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 5, 'The self-healing loop')
steps = [
    ('1 · PROBE', 'hit /health\n5 min, multi-region', ACCENT),
    ('2 · SCORE', 'uptime+latency\nerror rate → 0–100', CYAN),
    ('3 · DETECT', '3 failures in a row\n→ incident opens', AMBER),
    ('4 · ALERT', 'Slack / webhook\n+ alerts table', AMBER),
    ('5 · HEAL', 'GitHub deploy hook\nredeploys the cloud', GREEN),
    ('6 · VERIFY', 're-probe → close\nincident, log duration', GREEN),
]
x = Inches(0.7)
for label, body, color in steps:
    rect(s, x, Inches(2.1), Inches(1.92), Inches(2.5))
    text(s, x + Inches(0.15), Inches(2.3), Inches(1.62), Inches(0.4),
         label, size=14, bold=True, color=color)
    text(s, x + Inches(0.15), Inches(2.75), Inches(1.62), Inches(1.7),
         body, size=12, color=MUTED, leading=1.15)
    if label.split(' ')[0].strip() != '6':
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(1.94), Inches(3.15), Inches(0.18), Inches(0.4))
        a.fill.solid(); a.fill.fore_color.rgb = BORDER; a.line.fill.background(); a.shadow.inherit = False
    x += Inches(2.1)
rect(s, Inches(0.7), Inches(5.1), Inches(11.9), Inches(1.5), color=BG, line=BORDER)
text(s, Inches(1.0), Inches(5.3), Inches(11.3), Inches(1.2),
     [('Outcome:  ', {'size': 14, 'bold': True, 'color': GREEN}),
      ('a full outage becomes a short blip, handled with zero human intervention. Each heal is recorded in failover_events and the incident timeline for evidence.', {'size': 14, 'color': INK})])

# ============ SLIDE 6 — REAL-TIME ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 6, 'Real-time monitoring (v2.1)')
kpi(s, Inches(0.7), Inches(2.0), Inches(2.85), Inches(1.45), 'probe interval', '10 sec', 'live cloud probes', GREEN)
kpi(s, Inches(3.7), Inches(2.0), Inches(2.85), Inches(1.45), 'history', '5 min', 'cron → Supabase', ACCENT)
kpi(s, Inches(6.7), Inches(2.0), Inches(2.85), Inches(1.45), 'push channel', 'SSE', 'Server-Sent Events', CYAN)
kpi(s, Inches(9.7), Inches(2.0), Inches(2.85), Inches(1.45), 'sla target', '99.9%', 'Met / Near / At risk', AMBER)
bullet(s, Inches(0.7), Inches(3.85), Inches(11.9), Inches(2.9), [
    ('Server-Sent Events (/api/stream) push a fresh snapshot to every open dashboard every 10 s — the screen updates itself, no refresh.', {}),
    ('Live on-demand probes: the portal itself hits each cloud and shows status + latency as a pulsing strip.', {}),
    ('Live event feed: probes, routing decisions and incident changes scroll in a terminal-style log.', {}),
    ('Per-cloud extras: p50/p90/p95/p99, 24h latency trend arrows, sparklines, SLA vs 99.9%, embeddable SVG badges, CSV/JSON export, dark/light theme.', {}),
])

# ============ SLIDE 7 — INTELLIGENCE ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 7, 'Intelligence layer')
rows = [
    ('Anomaly detection', 'Isolation Forest (Python, CI) + rolling z-score detector (JS, live) — flags a cloud BEFORE it fails', ACCENT),
    ('Cost forecasting', '₹0/month operation vs ~$40/million single-cloud baseline; Cloud Run = $0.40/million beyond free tier', GREEN),
    ('Carbon-aware routing', 'grid intensity per region (live API or fallback averages) — greenest region wins ties', CYAN),
    ('Security & tenancy', 'Trivy + Checkov + OPA/Rego in CI · hashed API keys · per-key rate limiting · audit log', AMBER),
]
y = Inches(2.0)
for title, body, color in rows:
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.0))
    text(s, Inches(1.0), y + Inches(0.12), Inches(3.1), Inches(0.7),
         title, size=15, bold=True, color=color)
    text(s, Inches(4.3), y + Inches(0.18), Inches(8.0), Inches(0.7),
         body, size=13, color=MUTED, leading=1.05)
    y += Inches(1.15)
text(s, Inches(0.7), Inches(6.6), Inches(11.9), Inches(0.5),
     [('Also: OpenAPI/Swagger API, automated daily/weekly reports, 30 unit tests, Docker Compose one-command run.',
       {'size': 13, 'color': MUTED, 'align': PP_ALIGN.CENTER})])

# ============ SLIDE 8 — RESULTS ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 8, 'Live results (verified 2026-09-17)')
kpi(s, Inches(0.7), Inches(2.0), Inches(2.85), Inches(1.45), 'checks in db', '866', 'render 433 · vercel 433', ACCENT)
kpi(s, Inches(3.7), Inches(2.0), Inches(2.85), Inches(1.45), 'vercel score', '93', 'Excellent · 100% uptime', GREEN)
kpi(s, Inches(6.7), Inches(2.0), Inches(2.85), Inches(1.45), 'render score', '85', 'Good · ~98% uptime', AMBER)
kpi(s, Inches(9.7), Inches(2.0), Inches(2.85), Inches(1.45), 'total cost', '₹0/month', 'all three free tiers', GREEN)
rect(s, Inches(0.7), Inches(3.8), Inches(5.8), Inches(3.0))
text(s, Inches(1.0), Inches(3.95), Inches(5.2), Inches(0.4),
     'Evidence collected', size=15, bold=True)
bullet(s, Inches(1.0), Inches(4.4), Inches(5.2), Inches(2.3), [
    ('Router fails over to Vercel (score 93).', {'size': 13}),
    ('Anomaly detector flagged Render cold-start spikes.', {'size': 13}),
    ('30/30 tests green in GitHub Actions.', {'size': 13}),
    ('Incidents / alerts / audit tables recorded.', {'size': 13}),
])
rect(s, Inches(6.7), Inches(3.8), Inches(5.9), Inches(3.0))
text(s, Inches(7.0), Inches(3.95), Inches(5.2), Inches(0.4),
     'URLs to show', size=15, bold=True)
bullet(s, Inches(7.0), Inches(4.4), Inches(5.2), Inches(2.3), [
    ('Router: aegis-router.manav32478.workers.dev', {'size': 12}),
    ('Render: aegis-platform-pomf.onrender.com', {'size': 12}),
    ('Vercel: aegis-platform-lyart.vercel.app', {'size': 12}),
    ('Repo: github.com/Manav32478/aegis-platform-', {'size': 12}),
])

# ============ SLIDE 9 — FUTURE SCOPE ============
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 9, 'Future scope')
bullet(s, Inches(0.7), Inches(1.95), Inches(11.9), Inches(4.4), [
    ('Supabase Auth login + per-tenant dashboards (each API customer sees only their clouds).', {}),
    ('Move all Supabase writes behind an Edge Function with the service_role key — closes the demo-mode RLS gap.', {}),
    ('Synthetic transaction checks (login → dashboard → write) to catch “up but broken” clouds; TLS-certificate expiry monitoring.', {}),
    ('Auto-scaling rehearsal with k6 + a chaos button that kills a cloud on demand and films the recovery.', {}),
    ('ML-first failover: train on all tenants’ history, pre-warm the backup cloud before the failure; carbon-aware batch scheduling.', {}),
    ('Open-source release with a one-command installer; SOC2-style control mapping from the audit log.', {}),
])
text(s, Inches(0.7), Inches(6.3), Inches(11.9), Inches(0.6),
     [('Honest limits (say these): single-tenant demo auth · in-memory rate limiting · illustrative cost baseline · 2 of 3 clouds deployed.',
       {'size': 13, 'color': MUTED})])

# ============ SLIDE 10 — CLOSE ============
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.7), color=PANEL)
text(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(0.9),
     [('Questions?', {'size': 44, 'bold': True, 'align': PP_ALIGN.CENTER})])
text(s, Inches(0.9), Inches(3.95), Inches(11.5), Inches(0.5),
     [('Aegis — Self-Healing Multi-Cloud Platform · Manav Sarvaiya · ₹0/month',
       {'size': 15, 'color': MUTED, 'align': PP_ALIGN.CENTER})])

prs.save('/home/user/aegis-platform/docs/EVAL2_PRESENTATION.pptx')
print('saved', '/home/user/aegis-platform/docs/EVAL2_PRESENTATION.pptx')
